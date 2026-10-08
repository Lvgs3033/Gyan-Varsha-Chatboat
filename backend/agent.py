"""Gyan Varsha - Self-RAG agent: Gemini + local FastEmbed + FAISS + Tavily + Postgres memory.
Supports PDF, Excel/CSV (indexed for retrieval) and images (sent to Gemini vision)."""
from __future__ import annotations

from . import db  # keep first: sets up the libpq path on Windows before psycopg is imported

import csv
import io
import os
import shutil
import tempfile
from datetime import date
from pathlib import Path
from typing import Annotated, Any, Dict, List, Literal, Optional, TypedDict

from dotenv import load_dotenv
from langchain_community.document_loaders import PyPDFLoader
from langchain_community.embeddings import FastEmbedEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_core.messages import BaseMessage, HumanMessage, SystemMessage
from langchain_core.runnables import RunnableConfig
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langgraph.checkpoint.postgres import PostgresSaver
from langgraph.graph import END, START, StateGraph
from langgraph.graph.message import add_messages
from pydantic import BaseModel, Field

load_dotenv(override=True)

DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "faiss"
DATA_DIR.mkdir(parents=True, exist_ok=True)

# 1. LLM + embeddings --------------------------------------------------------------
MODEL = os.getenv("GEMINI_MODEL", "gemini-3.6-flash")
llm = ChatGoogleGenerativeAI(model=MODEL, temperature=0.3)          # answers
llm_strict = ChatGoogleGenerativeAI(model=MODEL, temperature=0)     # router / grader / verifier
SELF_CHECK = os.getenv("ENABLE_SELF_CHECK", "true").lower() == "true"
MAX_REWRITES = 1
embeddings = FastEmbedEmbeddings(model_name="BAAI/bge-small-en-v1.5")  # local, free

# 2. Document store (one FAISS index per chat, several files per chat) --------------
_RETRIEVERS: Dict[str, Any] = {}
TABLE_EXT = {".xlsx", ".xlsm", ".xls", ".csv"}
MAX_ROWS, ROWS_PER_CHUNK = 10000, 15


def _make_retriever(store):
    return store.as_retriever(search_kwargs={"k": 6})


def get_retriever(thread_id: Optional[str]):
    if not thread_id:
        return None
    if thread_id in _RETRIEVERS:
        return _RETRIEVERS[thread_id]
    path = DATA_DIR / thread_id
    if path.exists():
        store = FAISS.load_local(str(path), embeddings, allow_dangerous_deserialization=True)
        _RETRIEVERS[thread_id] = _make_retriever(store)
        return _RETRIEVERS[thread_id]
    return None


def _add_to_store(thread_id: str, chunks: list[Document]):
    path = DATA_DIR / thread_id
    if path.exists():
        store = FAISS.load_local(str(path), embeddings, allow_dangerous_deserialization=True)
        store.add_documents(chunks)
    else:
        store = FAISS.from_documents(chunks, embeddings)
    store.save_local(str(path))
    _RETRIEVERS[thread_id] = _make_retriever(store)


def _pdf_chunks(data: bytes, filename: str) -> list[Document]:
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as f:
        f.write(data)
        tmp = f.name
    try:
        pages = PyPDFLoader(tmp).load()
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass
    chunks = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200).split_documents(pages)
    out = []
    for c in chunks:
        if c.page_content.strip():
            c.metadata = {"file": filename, "loc": f"p.{c.metadata.get('page', 0) + 1}"}
            out.append(c)
    return out


def _load_sheets(data: bytes, filename: str) -> Dict[str, list]:
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        return {"Sheet1": list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))}
    if ext in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return {ws.title: [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    import xlrd  # .xls

    wb = xlrd.open_workbook(file_contents=data)
    return {sh.name: [[str(v) for v in sh.row_values(r)] for r in range(sh.nrows)] for sh in wb.sheets()}


def _num(v):
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _sheet_docs(filename: str, sheet: str, rows: list) -> list[Document]:
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return []
    header = [str(h).strip() or f"col{i + 1}" for i, h in enumerate(rows[0])]
    body = rows[1 : MAX_ROWS + 1]
    stats = []
    for j, h in enumerate(header):
        vals = [str(r[j]).strip() for r in body if j < len(r) and str(r[j]).strip()]
        nums = [n for n in (_num(v) for v in vals) if n is not None]
        if vals and len(nums) >= 0.8 * len(vals):
            stats.append(f"{h}: sum={sum(nums):g}, average={sum(nums) / len(nums):.4g}, min={min(nums):g}, max={max(nums):g}, count={len(nums)}")
    summary = f"Spreadsheet {filename}, sheet '{sheet}'. Columns: {', '.join(header)}. Data rows: {len(body)}."
    if stats:
        summary += " Numeric column statistics over all rows: " + "; ".join(stats) + "."
    docs = [Document(page_content=summary, metadata={"file": filename, "loc": f"{sheet} summary"})]
    for i in range(0, len(body), ROWS_PER_CHUNK):
        part = body[i : i + ROWS_PER_CHUNK]
        lines = [
            f"Row {i + k + 2}: " + " | ".join(f"{h}: {v}" for h, v in zip(header, r) if str(v).strip())
            for k, r in enumerate(part)
        ]
        docs.append(
            Document(
                page_content=f"Sheet '{sheet}'\n" + "\n".join(lines),
                metadata={"file": filename, "loc": f"{sheet} rows {i + 2}-{i + 1 + len(part)}"},
            )
        )
    return docs


def ingest_file(data: bytes, thread_id: str, filename: str) -> dict:
    """Index a PDF / Excel / CSV file into this chat's vector store."""
    if not data:
        raise ValueError("The uploaded file is empty.")
    existing = [d.get("filename") for d in (db.get_thread(thread_id) or {}).get("docs", [])]
    if filename in existing:
        raise ValueError(f"{filename} is already uploaded in this chat.")
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        kind, chunks = "pdf", _pdf_chunks(data, filename)
    elif ext in TABLE_EXT:
        kind, chunks = "sheet", []
        for name, rows in _load_sheets(data, filename).items():
            chunks += _sheet_docs(filename, name, rows)
    else:
        raise ValueError("Unsupported file type.")
    if not chunks:
        raise ValueError("No readable content found (scanned PDFs need OCR).")
    _add_to_store(thread_id, chunks)
    info = {"filename": filename, "kind": kind, "chunks": len(chunks)}
    db.add_doc(thread_id, info)
    return info


def drop_thread_data(thread_id: str):
    _RETRIEVERS.pop(thread_id, None)
    shutil.rmtree(DATA_DIR / thread_id, ignore_errors=True)


# 3. Helpers -----------------------------------------------------------------------
def text_of(content) -> str:
    """Gemini may return a string or a list of content parts."""
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        return "".join(p if isinstance(p, str) else p.get("text", "") for p in content if isinstance(p, (str, dict)))
    return ""


def has_image(msg) -> bool:
    return isinstance(msg.content, list) and any(isinstance(p, dict) and p.get("type") == "image_url" for p in msg.content)


def tavily_search(query: str) -> list[dict]:
    from tavily import TavilyClient

    key = os.getenv("TAVILY_API_KEY")
    if not key:
        raise RuntimeError("TAVILY_API_KEY is missing in .env")
    res = TavilyClient(api_key=key).search(query=query, max_results=5, search_depth="basic")
    return [{"text": r.get("content", ""), "title": r.get("title", ""), "url": r.get("url", "")} for r in res.get("results", [])]


def fmt_doc(ctx: list[dict]) -> str:
    return "\n\n".join(f"[{c['file']} - {c['loc']}] {c['text']}" for c in ctx)


def fmt_web(ctx: list[dict]) -> str:
    return "\n\n".join(f"[{i + 1}] {c['title']} ({c['url']})\n{c['text']}" for i, c in enumerate(ctx))


def recent_chat(messages, n=6) -> str:
    lines = []
    for m in messages[-n - 1 : -1]:
        who = "User" if isinstance(m, HumanMessage) else "Assistant"
        img = " [image attached]" if has_image(m) else ""
        lines.append(f"{who}: {text_of(m.content)[:400]}{img}")
    return "\n".join(lines) or "(no earlier messages)"


# 4. Structured outputs for the reflection steps -----------------------------------
class RouteDecision(BaseModel):
    route: Literal["direct", "document", "web"]
    standalone_question: str = Field(description="The user's latest question rewritten so it makes sense without the chat history.")


class GradeResult(BaseModel):
    relevant_ids: List[int] = Field(description="Ids of excerpts that help answer the question.")
    answerable: bool = Field(description="True if the relevant excerpts contain enough to answer the question.")


class Verdict(BaseModel):
    grounded: bool = Field(description="True if every claim in the answer is supported by the context.")


# 5. State -------------------------------------------------------------------------
class ChatState(TypedDict, total=False):
    messages: Annotated[list[BaseMessage], add_messages]
    question: str
    route: str
    source: str
    contexts: list
    fallback: bool
    good: bool
    rewrites: int
    meta: dict
    grounded: Optional[bool]


# 6. Nodes -------------------------------------------------------------------------
def route_node(state: ChatState, config: RunnableConfig):
    last = state["messages"][-1]
    text = text_of(last.content)
    if has_image(last):  # the user attached an image: Gemini vision answers directly
        return {"route": "direct", "question": text, "source": "direct"}
    thread_id = config["configurable"]["thread_id"]
    has_doc = get_retriever(thread_id) is not None
    earlier_img = any(has_image(m) for m in state["messages"][-7:-1])
    common_direct = "- direct: general knowledge, explanations, coding, maths, writing, small talk, stable facts"
    web = "- web: needs up-to-date or real-time information (news, prices, scores, weather, recent releases) or facts you are unsure about."
    if has_doc:
        options = (
            "- document: about the contents of the user's uploaded files (PDF, Excel or CSV): summaries, 'in the document/sheet', "
            "or follow-ups to earlier file answers. If ambiguous and files exist, choose document.\n"
            f"{common_direct} that do not need the files.\n{web}"
        )
    else:
        options = f"{common_direct}.\n{web}\n(No files are uploaded, so never choose document.)"
    if earlier_img:
        options += "\n(An image was shared earlier in this chat. Questions about that image: choose direct.)"
    prompt = (
        f"Today is {date.today()}. Decide how to answer the user's latest message.\n{options}\n\n"
        f"Chat so far:\n{recent_chat(state['messages'])}\n\nLatest message: {text}"
    )
    try:
        d = llm_strict.with_structured_output(RouteDecision).invoke(prompt)
        route, q = d.route, d.standalone_question or text
    except Exception:
        route, q = ("document" if has_doc else "direct"), text
    if route == "document" and not has_doc:
        route = "direct"
    return {"route": route, "question": q, "source": route}


def retrieve_node(state: ChatState, config: RunnableConfig):
    retriever = get_retriever(config["configurable"]["thread_id"])
    docs = retriever.invoke(state["question"]) if retriever else []
    return {
        "contexts": [
            {
                "text": d.page_content,
                "file": d.metadata.get("file", "document"),
                "loc": d.metadata.get("loc") or f"p.{d.metadata.get('page', 0) + 1}",
            }
            for d in docs
        ]
    }


def grade_node(state: ChatState):
    ctx = state.get("contexts", [])
    if not ctx:
        return {"good": False, "contexts": []}
    listing = "\n\n".join(f"[{i}] ({c['file']} - {c['loc']}) {c['text']}" for i, c in enumerate(ctx))
    prompt = f"Question: {state['question']}\n\nExcerpts from the user's files:\n{listing}\n\nWhich excerpts are relevant, and can the question be answered from them?"
    try:
        g = llm_strict.with_structured_output(GradeResult).invoke(prompt)
        keep = [ctx[i] for i in g.relevant_ids if 0 <= i < len(ctx)]
        return {"good": bool(g.answerable and keep), "contexts": keep}
    except Exception:
        return {"good": True}


def rewrite_node(state: ChatState):
    prompt = (
        "Rewrite this question as a short keyword-rich search query for finding the answer inside a document or spreadsheet. "
        f"Reply with the query only.\n\nQuestion: {state['question']}"
    )
    q = text_of(llm_strict.invoke(prompt).content).strip() or state["question"]
    return {"question": q, "rewrites": state.get("rewrites", 0) + 1}


def web_node(state: ChatState):
    fallback = state.get("route") == "document"
    try:
        results = tavily_search(state["question"])
    except Exception as e:
        results = []
        print("Tavily error:", e)
    return {"contexts": results, "source": "web", "fallback": fallback}


def generate_node(state: ChatState, config: RunnableConfig):
    source, ctx = state.get("source", "direct"), state.get("contexts", [])
    if source == "direct":
        system = (
            "You are Gyan Varsha, a helpful assistant. Answer from your own knowledge (and any image the user attached), "
            "clearly, in Markdown. If the user asks about an uploaded file but none is available, tell them to upload it. "
            "If you are unsure or the topic may have changed recently, say so."
        )
    elif source == "document":
        system = (
            "You are Gyan Varsha. Answer using ONLY the excerpts from the user's uploaded files below. "
            "Cite file and location like (report.pdf, p.3) or (sales.xlsx, Sheet1 rows 2-16). Use Markdown.\n\nEXCERPTS:\n" + fmt_doc(ctx)
        )
    else:
        intro = (
            "The uploaded files did not contain the answer: say that in one short sentence first, then answer from the web results. "
            if state.get("fallback")
            else ""
        )
        if ctx:
            system = f"You are Gyan Varsha. {intro}Answer using the web results below. Cite sources like [1]. Use Markdown.\n\nWEB RESULTS:\n{fmt_web(ctx)}"
        else:
            system = (
                f"You are Gyan Varsha. {intro}Live web search returned nothing (it may be unavailable). Say so briefly, "
                "then give your best answer from your own knowledge and flag uncertainty."
            )
    resp = llm.invoke([SystemMessage(content=system), *state["messages"][-10:]], config=config)
    if source == "document":
        seen, sources = set(), []
        for c in ctx:
            key = (c["file"], c["loc"])
            if key not in seen:
                seen.add(key)
                sources.append({"file": c["file"], "loc": c["loc"]})
        sources = sources[:6]
    elif source == "web":
        sources = [{"title": c["title"], "url": c["url"]} for c in ctx]
    else:
        sources = []
    meta = {"source": source, "fallback": bool(state.get("fallback")), "sources": sources}
    resp.additional_kwargs["rag"] = meta
    return {"messages": [resp], "meta": meta}


def verify_node(state: ChatState):
    answer = text_of(state["messages"][-1].content)
    ctx = fmt_doc(state["contexts"]) if state["source"] == "document" else fmt_web(state["contexts"])
    try:
        v = llm_strict.with_structured_output(Verdict).invoke(f"Context:\n{ctx}\n\nAnswer:\n{answer}\n\nIs the answer supported by the context?")
        return {"grounded": v.grounded}
    except Exception:
        return {"grounded": None}


# 7. Graph (Self-RAG) --------------------------------------------------------------
def after_route(state: ChatState):
    return {"direct": "generate", "document": "retrieve", "web": "web_search"}[state["route"]]


def after_grade(state: ChatState):
    if state.get("good"):
        return "generate"
    return "rewrite" if state.get("rewrites", 0) < MAX_REWRITES else "web_search"


def after_generate(state: ChatState):
    return "verify" if SELF_CHECK and state.get("source") in ("document", "web") and state.get("contexts") else END


graph = StateGraph(ChatState)
for name, fn in [("route", route_node), ("retrieve", retrieve_node), ("grade", grade_node), ("rewrite", rewrite_node),
                 ("web_search", web_node), ("generate", generate_node), ("verify", verify_node)]:
    graph.add_node(name, fn)
graph.add_edge(START, "route")
graph.add_conditional_edges("route", after_route, ["generate", "retrieve", "web_search"])
graph.add_edge("retrieve", "grade")
graph.add_conditional_edges("grade", after_grade, ["generate", "rewrite", "web_search"])
graph.add_edge("rewrite", "retrieve")
graph.add_edge("web_search", "generate")
graph.add_conditional_edges("generate", after_generate, ["verify", END])
graph.add_edge("verify", END)

db.init_db()
checkpointer = PostgresSaver(db.pool)
checkpointer.setup()
chatbot = graph.compile(checkpointer=checkpointer)
